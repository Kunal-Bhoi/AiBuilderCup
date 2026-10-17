"""File handling for ingest (§11.3): magic-byte sniffing, conversion to PDF, page rendering.

Licensing: pypdfium2 (Apache-2.0/BSD), pypdf (BSD), Pillow (HPND), openpyxl (MIT),
reportlab (BSD). PyMuPDF is deliberately not used because it is AGPL.
Office conversion uses openpyxl + reportlab rather than LibreOffice so the service builds
with Cloud Run source deploys (buildpacks, no Dockerfile)."""

import io
import threading
import zipfile
from dataclasses import dataclass, field
from enum import StrEnum

from PIL import Image, ImageOps

from app.errors import ErrorCode

try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
except Exception:  # pragma: no cover - optional codec
    pass

# PDFium is not thread-safe; serialise all access to it within the process.
_pdfium_lock = threading.Lock()


class Kind(StrEnum):
    PDF = "PDF"
    PNG = "PNG"
    JPEG = "JPEG"
    TIFF = "TIFF"
    WEBP = "WEBP"
    HEIC = "HEIC"
    XLSX = "XLSX"
    UNKNOWN = "UNKNOWN"


IMAGE_KINDS = {Kind.PNG, Kind.JPEG, Kind.TIFF, Kind.WEBP, Kind.HEIC}
SUPPORTED_KINDS = IMAGE_KINDS | {Kind.PDF, Kind.XLSX}

# Declared content types accepted at upload-URL time. The declared type is only a first
# filter; ingest re-checks the real bytes, because a declared content type is a lie by default.
ACCEPTED_CONTENT_TYPES = {
    "application/pdf",
    "image/png",
    "image/jpeg",
    "image/jpg",
    "image/tiff",
    "image/webp",
    "image/heic",
    "image/heif",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",
}


def sniff(blob: bytes) -> Kind:
    head = blob[:16]
    if head.startswith(b"%PDF") or b"%PDF" in blob[:1024]:
        return Kind.PDF
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return Kind.PNG
    if head.startswith(b"\xff\xd8\xff"):
        return Kind.JPEG
    if head.startswith((b"II*\x00", b"MM\x00*")):
        return Kind.TIFF
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return Kind.WEBP
    if head[4:8] == b"ftyp" and head[8:12] in (b"heic", b"heix", b"mif1", b"heif", b"hevc", b"msf1"):
        return Kind.HEIC
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(blob)) as z:
                if any(n.startswith("xl/") for n in z.namelist()):
                    return Kind.XLSX
        except zipfile.BadZipFile:
            return Kind.UNKNOWN
    return Kind.UNKNOWN


class IngestFailure(Exception):
    def __init__(self, code: ErrorCode, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class RenderedPage:
    page: int
    png: bytes
    width_px: int
    height_px: int


@dataclass
class IngestOutput:
    kind: Kind
    page_count: int
    text_layer: str | None
    pages: list[RenderedPage] = field(default_factory=list)


def image_to_pdf(blob: bytes) -> bytes:
    with Image.open(io.BytesIO(blob)) as im:
        frames = []
        for i in range(getattr(im, "n_frames", 1)):
            im.seek(i)
            frames.append(ImageOps.exif_transpose(im.copy()).convert("RGB"))
    out = io.BytesIO()
    frames[0].save(out, format="PDF", save_all=True, append_images=frames[1:], resolution=200)
    return out.getvalue()


def xlsx_to_pdf(blob: bytes) -> bytes:
    from openpyxl import load_workbook
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

    wb = load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    styles = getSampleStyleSheet()
    story = []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c.strip() for c in r)][:500]
        story.append(Paragraph(ws.title, styles["Heading2"]))
        if rows:
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            t = Table(rows, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 7)]))
            story.append(t)
        story.append(PageBreak())
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=landscape(A4)).build(story or [Paragraph("(empty workbook)", styles["Normal"])])
    return out.getvalue()


def _is_blank(img: Image.Image) -> bool:
    lo, hi = img.convert("L").getextrema()
    return hi - lo < 8


def _downscale(img: Image.Image, max_px: int) -> Image.Image:
    longest = max(img.size)
    if longest <= max_px:
        return img
    scale = max_px / longest
    return img.resize((max(1, int(img.width * scale)), max(1, int(img.height * scale))), Image.LANCZOS)


def ingest_bytes(blob: bytes, *, dpi: int, max_px: int, max_pages: int) -> IngestOutput:
    """Validate, convert and render. Raises IngestFailure with a §7 code on a permanent failure.
    Synchronous and CPU-bound: call through asyncio.to_thread."""
    import pypdfium2 as pdfium
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    if not blob:
        raise IngestFailure(ErrorCode.DOCUMENT_CORRUPT, "The file is empty.")

    kind = sniff(blob)
    if kind not in SUPPORTED_KINDS:
        raise IngestFailure(ErrorCode.UNSUPPORTED_FILE_TYPE, "The file type is not supported.")

    try:
        if kind in IMAGE_KINDS:
            blob = image_to_pdf(blob)
        elif kind == Kind.XLSX:
            blob = xlsx_to_pdf(blob)
    except Exception:
        raise IngestFailure(ErrorCode.DOCUMENT_CORRUPT, "The file could not be decoded.")

    try:
        reader = PdfReader(io.BytesIO(blob))
        if reader.is_encrypted and not reader.decrypt(""):
            raise IngestFailure(ErrorCode.DOCUMENT_PASSWORD_PROTECTED, "This PDF is password protected and cannot be read.")
    except IngestFailure:
        raise
    except (PdfReadError, Exception) as exc:
        if "password" in str(exc).lower() or "encrypt" in str(exc).lower():
            raise IngestFailure(ErrorCode.DOCUMENT_PASSWORD_PROTECTED, "This PDF is password protected and cannot be read.")
        raise IngestFailure(ErrorCode.DOCUMENT_CORRUPT, "The PDF is corrupt and cannot be read.")

    with _pdfium_lock:
        try:
            pdf = pdfium.PdfDocument(blob)
        except pdfium.PdfiumError as exc:
            if "password" in str(exc).lower():
                raise IngestFailure(ErrorCode.DOCUMENT_PASSWORD_PROTECTED, "This PDF is password protected and cannot be read.")
            raise IngestFailure(ErrorCode.DOCUMENT_CORRUPT, "The PDF is corrupt and cannot be read.")
        try:
            n = len(pdf)
            if n == 0:
                raise IngestFailure(ErrorCode.NO_TEXT_OR_IMAGE_CONTENT, "The document has no pages.")
            if n > max_pages:
                raise IngestFailure(
                    ErrorCode.DOCUMENT_PAGE_LIMIT_EXCEEDED, f"The document has {n} pages; the limit is {max_pages}."
                )
            texts: list[str] = []
            pages: list[RenderedPage] = []
            blank = 0
            for i in range(n):
                page = pdf[i]
                textpage = page.get_textpage()
                texts.append(textpage.get_text_range())
                textpage.close()
                img = page.render(scale=dpi / 72, rotation=0).to_pil().convert("RGB")
                page.close()
                if _is_blank(img):
                    blank += 1
                img = _downscale(img, max_px)
                buf = io.BytesIO()
                img.save(buf, format="PNG", optimize=True)
                pages.append(RenderedPage(page=i + 1, png=buf.getvalue(), width_px=img.width, height_px=img.height))
        finally:
            pdf.close()

    text_layer = "\n".join(texts)
    has_text = len(text_layer.strip()) > 50
    if not has_text and blank == len(pages):
        raise IngestFailure(ErrorCode.NO_TEXT_OR_IMAGE_CONTENT, "Every page is blank.")
    return IngestOutput(kind=kind, page_count=len(pages), text_layer=text_layer if has_text else None, pages=pages)


def crop_png(png: bytes, bbox: tuple[float, float, float, float], margin: float = 0.0) -> bytes:
    with Image.open(io.BytesIO(png)) as im:
        w, h = im.size
        x0, y0, x1, y1 = bbox
        mx, my = (x1 - x0) * margin, (y1 - y0) * margin
        box = (
            max(0, int((x0 - mx) * w)),
            max(0, int((y0 - my) * h)),
            min(w, max(int((x1 + mx) * w), int((x0 - mx) * w) + 1)),
            min(h, max(int((y1 + my) * h), int((y0 - my) * h) + 1)),
        )
        out = io.BytesIO()
        im.crop(box).save(out, format="PNG")
        return out.getvalue()
